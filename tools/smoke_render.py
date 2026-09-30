"""CI 검증: 렌더러(모든 요소) → PPTX → 역변환(importer) → 원본 템플릿에 장표 교체."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pptx import Presentation  # noqa: E402

from studio.importer import import_slide  # noqa: E402
from studio.llm import _mock_layout, placeholder_png  # noqa: E402
from studio.render_pptx import Renderer  # noqa: E402

work = Path("smoke_out")
work.mkdir(exist_ok=True)
img = work / "logo.png"
img.write_bytes(placeholder_png(320, 180))

layout = _mock_layout(False)
layout["elements"].append({"id": "e20", "type": "image", "x": 70, "y": 88, "w": 10, "h": 10, "source": "user:logo.png"})
layout["elements"].append({"id": "e21", "type": "image", "x": 82, "y": 88, "w": 10, "h": 10, "source": "gen:missing"})
layout["elements"].append({"id": "e22", "type": "ellipse", "x": 50, "y": 88, "w": 3, "h": 5, "fill": "#FF6B2C"})

# 1) 렌더
r = Renderer(palette={"#FFFFFF", "#111111", "#747474", "#E7E7E4", "#FF6B2C", "#C9C9C6"})
_, resolved = r.add_slide(layout, {"user:logo.png": img})
out = r.save(work / "slide.pptx")
print("warnings:", r.warnings)
assert not [w for w in r.warnings if "렌더 실패" in w], r.warnings
prs = Presentation(str(out))
names = [s.name for s in prs.slides[0].shapes]
print("shapes:", names)
assert len(names) == len(layout["elements"]), (len(names), len(layout["elements"]))
assert "e20" in resolved

# 2) 역변환 (PowerPoint 에서 고친 파일을 다시 읽는 경로)
back, warns = import_slide(out, work / "media", "smoke_out/media")
print("import warnings:", warns)
ids = {e["id"]: e for e in back["elements"]}
for want in ("e1", "e2", "e3", "e4", "e7", "e10", "e20", "e22"):
    assert want in ids, (want, list(ids))
assert ids["e1"]["type"] == "text" and "데모" in ids["e1"]["text"], ids["e1"]
assert ids["e4"]["type"] == "donut" and ids["e4"]["values"] == [1.0, 4.0], ids["e4"]
assert ids["e7"]["type"] == "bar" and ids["e7"]["horizontal"] is True, ids["e7"]
assert ids["e10"]["type"] == "line" and ids["e10"]["dash"] and ids["e10"]["arrow"], ids["e10"]
assert ids["e20"]["type"] == "image", ids["e20"]
assert abs(ids["e1"]["x"] - 6) < 0.5, ids["e1"]
print("import OK:", len(back["elements"]), "elements")

# 3) 다시 렌더 (역변환 결과가 렌더 가능한지)
r2 = Renderer()
r2.add_slide(back, {})
r2.save(work / "slide-roundtrip.pptx")

# 4) 원본 템플릿(3장)에서 2번만 교체
tpl = Presentation()
for i in range(3):
    s = tpl.slides.add_slide(tpl.slide_layouts[5])
    s.shapes.title.text = f"원본 {i + 1}"
tpl.save(str(work / "template.pptx"))
r3 = Renderer(template=work / "template.pptx", keep_slides=True)
r3.add_slide(layout, {})
r3.replace_slides({2: r3.last_sid})
deck = r3.save(work / "deck.pptx")
d = Presentation(str(deck))
assert len(d.slides) == 3, len(d.slides)
titles = [next((sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text), "") for s in d.slides]
print("deck:", titles)
assert titles[0] == "원본 1" and titles[2] == "원본 3" and "데모" in titles[1], titles

# 5) 템플릿 기반 단일 장표 (기존 슬라이드 제거)
r4 = Renderer(template=work / "template.pptx")
r4.add_slide(layout, {})
assert len(Presentation(str(r4.save(work / "single.pptx"))).slides) == 1
print("OK")
