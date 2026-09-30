"""사용자가 PowerPoint 에서 직접 고친 PPTX → 레이아웃 JSON 으로 역변환.

이렇게 하면 AI 가 다음 수정을 할 때 사용자가 손본 내용을 덮어쓰지 않고 그 위에서 고친다.
"""
from __future__ import annotations

from pathlib import Path


def _hex(color_format) -> str | None:
    try:
        if color_format.type is not None and color_format.rgb is not None:
            return "#" + str(color_format.rgb)
    except Exception:
        pass
    return None


def import_slide(pptx_path: Path, media_dir: Path, media_prefix: str, slide_index: int = 0) -> tuple[dict, list[str]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.dml import MSO_FILL
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn

    prs = Presentation(str(pptx_path))
    if slide_index >= len(prs.slides):
        raise ValueError("PPTX에 장표가 없습니다")
    slide = prs.slides[slide_index]
    W, H = prs.slide_width, prs.slide_height
    warns: list[str] = []
    els: list[dict] = []
    media_dir.mkdir(parents=True, exist_ok=True)
    counter = [0]

    def pct(v, total):
        return round(float(v or 0) / total * 100, 2)

    def box(shp):
        return {"x": pct(shp.left, W), "y": pct(shp.top, H), "w": pct(shp.width, W), "h": pct(shp.height, H)}

    def eid(shp):
        name = shp.name or ""
        if name.startswith("el:"):
            return name[3:]
        counter[0] += 1
        return f"m{counter[0]}"

    def text_el(shp, base):
        tf = shp.text_frame
        runs, size, bold, color, align = [], None, False, None, None
        for pi, p in enumerate(tf.paragraphs):
            if pi > 0:
                runs.append({"text": "\n"})
            if align is None and p.alignment is not None:
                align = {PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right"}.get(p.alignment, "left")
            for r in p.runs:
                item = {"text": r.text}
                if r.font.size:
                    item["size"] = round(r.font.size.pt, 1)
                    size = size or item["size"]
                if r.font.bold:
                    item["bold"] = True
                    bold = True
                c = _hex(r.font.color)
                if c:
                    item["color"] = c
                    color = color or c
                runs.append(item)
        # 연속된 "\n" run 을 앞 run 에 합치기
        merged = []
        for r in runs:
            if r["text"] == "\n" and merged:
                merged[-1] = {**merged[-1], "text": merged[-1]["text"] + "\n"}
            else:
                merged.append(r)
        el = {**base, "type": "text", "size": size or 14, "color": color or "#111111", "bold": bold,
              "align": align or "left"}
        if len({(r.get("color"), r.get("bold"), r.get("size")) for r in merged}) <= 1:
            el["text"] = "".join(r["text"] for r in merged)
        else:
            el["runs"] = merged
        return el

    def walk(shapes):
        for shp in shapes:
            try:
                st = shp.shape_type
                base = {"id": eid(shp)}
                if st == MSO_SHAPE_TYPE.GROUP:
                    walk(shp.shapes)
                    continue
                if st == MSO_SHAPE_TYPE.PICTURE:
                    counter[0] += 1
                    img = shp.image
                    f = media_dir / f"{base['id']}-{counter[0]}.{img.ext}"
                    f.write_bytes(img.blob)
                    els.append({**base, **box(shp), "type": "image", "source": f"file:{media_prefix}/{f.name}",
                                "fit": "stretch"})
                    continue
                if getattr(shp, "has_chart", False) and shp.has_chart:
                    ch = shp.chart
                    plot = ch.plots[0]
                    ser = plot.series[0]
                    cats = [str(c) for c in plot.categories]
                    vals = [float(v or 0) for v in ser.values]
                    colors = []
                    for pt_i in range(len(vals)):
                        try:
                            colors.append(_hex(ser.points[pt_i].format.fill.fore_color) or "#C9C9C6")
                        except Exception:
                            colors.append("#C9C9C6")
                    if ch.chart_type == XL_CHART_TYPE.DOUGHNUT:
                        hole = 0.68
                        for hs in ch._chartSpace.iter(qn("c:holeSize")):
                            hole = int(hs.get("val", "68")) / 100
                        els.append({**base, **box(shp), "type": "donut", "values": vals, "labels": cats,
                                    "colors": colors, "hole": hole})
                    else:
                        els.append({**base, **box(shp), "type": "bar", "categories": cats, "values": vals,
                                    "colors": colors,
                                    "horizontal": ch.chart_type in (XL_CHART_TYPE.BAR_CLUSTERED,
                                                                    XL_CHART_TYPE.BAR_STACKED)})
                    continue
                if st in (MSO_SHAPE_TYPE.LINE,) or shp.__class__.__name__ == "Connector":
                    ln = shp.line
                    xml = shp._element.xml
                    els.append({**base, "type": "line",
                                "x1": pct(shp.begin_x, W), "y1": pct(shp.begin_y, H),
                                "x2": pct(shp.end_x, W), "y2": pct(shp.end_y, H),
                                "color": _hex(ln.color) or "#111111",
                                "width": round(ln.width.pt, 2) if ln.width else 1,
                                "dash": ln.dash_style is not None, "arrow": "tailEnd" in xml or "headEnd" in xml})
                    continue
                if st == MSO_SHAPE_TYPE.TEXT_BOX or (st == MSO_SHAPE_TYPE.PLACEHOLDER and shp.has_text_frame):
                    els.append(text_el(shp, {**base, **box(shp)}))
                    continue
                if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
                    fill = None
                    try:
                        if shp.fill.type == MSO_FILL.SOLID:
                            fill = _hex(shp.fill.fore_color)
                    except Exception:
                        pass
                    line = None
                    try:
                        if shp.line.fill.type == MSO_FILL.SOLID:
                            line = _hex(shp.line.color)
                    except Exception:
                        pass
                    ast = shp.auto_shape_type
                    if not fill and not line and shp.has_text_frame and shp.text_frame.text.strip():
                        els.append(text_el(shp, {**base, **box(shp)}))
                        continue
                    kind = "ellipse" if ast == MSO_SHAPE.OVAL else "rect"
                    el = {**base, **box(shp), "type": kind, "fill": fill, "line": line,
                          "line_w": round(shp.line.width.pt, 2) if shp.line.width else 0.75}
                    if ast == MSO_SHAPE.ROUNDED_RECTANGLE:
                        try:
                            el["radius"] = round(float(shp.adjustments[0]), 3)
                        except Exception:
                            el["radius"] = 0.1
                    if shp.has_text_frame and shp.text_frame.text.strip():
                        t = text_el(shp, {})
                        el.update(text=shp.text_frame.text, size=t["size"], color=t["color"], bold=t["bold"])
                    els.append(el)
                    continue
                warns.append(f"읽지 못한 도형 건너뜀: {shp.name}")
            except Exception as e:
                warns.append(f"도형 읽기 실패 {getattr(shp, 'name', '?')}: {e}")

    walk(slide.shapes)
    bg = "#FFFFFF"
    try:
        if slide.background.fill.type == MSO_FILL.SOLID:
            bg = _hex(slide.background.fill.fore_color) or bg
    except Exception:
        pass
    notes = ""
    try:
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
    except Exception:
        pass
    return {"background": bg, "elements": els, "notes": notes}, warns
