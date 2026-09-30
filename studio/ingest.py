"""원본 자료(PDF / PPTX / 이미지)를 장표별 이미지 + 텍스트로 분해."""
from __future__ import annotations

import platform
import shutil
import subprocess
import tempfile
from pathlib import Path

IMG_EXT = {".png", ".jpg", ".jpeg", ".webp"}


def pptx_to_pdf(src: Path, out_dir: Path) -> Path | None:
    """PowerPoint(Windows) → LibreOffice 순으로 PDF 변환 시도."""
    pdf = out_dir / (src.stem + ".pdf")
    if platform.system() == "Windows":
        ps = (
            "$ErrorActionPreference='Stop';"
            "$app=New-Object -ComObject PowerPoint.Application;"
            f"$p=$app.Presentations.Open('{src.resolve()}',$true,$false,$false);"
            f"$p.SaveAs('{pdf.resolve()}',32);$p.Close();$app.Quit()"
        )
        try:
            subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, timeout=300,
                           capture_output=True)
            if pdf.exists():
                return pdf
        except Exception:
            pass
    for exe in ("soffice", "libreoffice",
                "/Applications/LibreOffice.app/Contents/MacOS/soffice",
                r"C:\Program Files\LibreOffice\program\soffice.exe"):
        if shutil.which(exe) or Path(exe).exists():
            try:
                subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(src)],
                               check=True, timeout=300, capture_output=True)
                if pdf.exists():
                    return pdf
            except Exception:
                continue
    return None


def pptx_texts(src: Path) -> list[str]:
    try:
        from pptx import Presentation
    except ImportError:
        return []
    texts = []
    for slide in Presentation(str(src)).slides:
        chunks = []
        for shp in slide.shapes:
            if shp.has_text_frame:
                chunks.append(shp.text_frame.text)
            if getattr(shp, "has_table", False) and shp.has_table:
                for row in shp.table.rows:
                    chunks.append(" | ".join(c.text for c in row.cells))
        texts.append("\n".join(c for c in chunks if c.strip()))
    return texts


def pdf_pages(pdf: Path, pages_dir: Path, start: int) -> list[dict]:
    import fitz  # PyMuPDF

    out = []
    doc = fitz.open(str(pdf))
    for i, page in enumerate(doc):
        no = start + i
        zoom = 1600 / max(page.rect.width, 1)
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        img = pages_dir / f"{no:02d}.png"
        pix.save(str(img))
        out.append({"no": no, "image": f"pages/{img.name}", "text": page.get_text("text").strip()})
    return out


def ingest(files: list[Path], project_dir: Path) -> tuple[list[dict], str]:
    """반환: (pages, 경고 메시지)."""
    pages_dir = project_dir / "pages"
    pages_dir.mkdir(exist_ok=True)
    pages: list[dict] = []
    notes = []
    for f in sorted(files, key=lambda p: p.name):
        ext = f.suffix.lower()
        start = len(pages) + 1
        if ext in IMG_EXT:
            dst = pages_dir / f"{start:02d}{ext}"
            shutil.copy(f, dst)
            pages.append({"no": start, "image": f"pages/{dst.name}", "text": ""})
        elif ext == ".pdf":
            pages += pdf_pages(f, pages_dir, start)
        elif ext == ".pptx":
            texts = pptx_texts(f)
            with tempfile.TemporaryDirectory() as td:
                pdf = pptx_to_pdf(f, Path(td))
                if pdf:
                    got = pdf_pages(pdf, pages_dir, start)
                    for i, g in enumerate(got):
                        if i < len(texts) and texts[i]:
                            g["text"] = texts[i]
                    pages += got
                else:
                    notes.append("PPTX를 이미지로 바꿀 PowerPoint/LibreOffice가 없어 텍스트만 읽었습니다. "
                                 "장표 이미지가 필요하면 PDF로 내보내서 올려주세요.")
                    for i, t in enumerate(texts):
                        pages.append({"no": start + i, "image": "", "text": t})
        else:
            notes.append(f"지원하지 않는 형식: {f.name}")
    return pages, "\n".join(notes)
