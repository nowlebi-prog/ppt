"""레이아웃 JSON → 편집 가능한 PPTX.

좌표는 모두 슬라이드 대비 퍼센트(0~100).
허용 요소: text, rect, ellipse, line, image, donut, bar
→ 손글씨/3D 장식 같은 요소는 '부품'이 없으므로 PPTX에 들어갈 수 없다.
각 도형 이름에 "el:<id>" 를 넣어 두어, 사용자가 PowerPoint 에서 고친 파일을 다시 읽을 수 있다.
"""
from __future__ import annotations

import math
import re
from pathlib import Path

HEX = re.compile(r"^#?[0-9A-Fa-f]{6}$")
DEFAULT_W_PT, DEFAULT_H_PT = 960.0, 540.0  # 13.333 x 7.5 inch


def _rgb(hexstr, default="111111"):
    from pptx.dml.color import RGBColor
    h = hexstr if isinstance(hexstr, str) and HEX.match(hexstr) else default
    return RGBColor.from_string(h.lstrip("#").upper())


def _set_font(run, name: str):
    from pptx.oxml.ns import qn
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", name)


def text_of(el: dict) -> str:
    if el.get("runs"):
        return "".join(str(r.get("text", "")) for r in el["runs"])
    return str(el.get("text", ""))


def overflow_warnings(layout: dict, slide_w_pt: float = DEFAULT_W_PT, slide_h_pt: float = DEFAULT_H_PT) -> list[str]:
    """글자 수 기반으로 텍스트가 상자를 넘치는지 대략 추정 (한글 1자 ≈ 글자크기 폭)."""
    out = []
    for el in layout.get("elements", []):
        if el.get("type") != "text":
            continue
        size = float(el.get("size", 14))
        w_pt = slide_w_pt * float(el.get("w", 10)) / 100
        h_pt = slide_h_pt * float(el.get("h", 5)) / 100
        lines = 0
        for para in text_of(el).split("\n"):
            width = sum(size * (1.0 if ord(c) > 0x2E80 else 0.55) for c in para)
            lines += max(1, math.ceil(width / max(w_pt, 1)))
        need = lines * size * float(el.get("line_spacing", 1.2))
        if need > h_pt * 1.25:
            out.append(f"{el.get('id', '?')}: 텍스트가 상자를 넘칠 수 있음 (약 {lines}줄, \"{text_of(el)[:18]}…\")")
    return out


def normalize(layout: dict) -> dict:
    """id 부여, 숫자 보정."""
    els = layout.get("elements") or []
    used = {str(e.get("id")) for e in els if e.get("id")}
    n = 1
    for e in els:
        if not e.get("id"):
            while f"e{n}" in used:
                n += 1
            e["id"] = f"e{n}"
            used.add(e["id"])
        for k in ("x", "y", "w", "h", "x1", "y1", "x2", "y2"):
            if k in e:
                try:
                    e[k] = max(0.0, min(100.0, float(e[k])))
                except (TypeError, ValueError):
                    e[k] = 0.0
    layout["elements"] = els
    layout.setdefault("background", "#FFFFFF")
    return layout


class Renderer:
    def __init__(self, template: Path | None = None, keep_slides: bool = False, font: str = "Pretendard",
                 palette: set[str] | None = None, min_pt: float = 8):
        from pptx import Presentation
        from pptx.util import Pt
        if template and Path(template).exists():
            self.prs = Presentation(str(template))
            if not keep_slides:
                self._drop_all()
        else:
            self.prs = Presentation()
            self.prs.slide_width = Pt(DEFAULT_W_PT)
            self.prs.slide_height = Pt(DEFAULT_H_PT)
        self.font = font
        self.palette = {p.upper().lstrip("#") for p in (palette or set())}
        self.min_pt = min_pt
        self.warnings: list[str] = []

    # ---- 템플릿 처리
    def _drop_all(self):
        lst = self.prs.slides._sldIdLst
        for sid in list(lst):
            self.prs.part.drop_rel(sid.rId)
            lst.remove(sid)

    def _blank_layout(self):
        layouts = list(self.prs.slide_layouts)
        for lo in layouts:
            if lo.name and lo.name.lower() in ("blank", "빈 화면", "빈화면", "empty"):
                return lo
        return min(layouts, key=lambda lo: len(lo.placeholders))

    def replace_slides(self, new_by_no: dict):
        """keep_slides 모드: 원본 N번 슬라이드를 새로 만든 슬라이드(sldId)로 교체. 나머지 원본은 그대로."""
        lst = self.prs.slides._sldIdLst
        ids = list(lst)
        new_ids = dict(new_by_no)
        originals = [s for s in ids if s not in new_ids.values()]
        order = []
        for i, sid in enumerate(originals, start=1):
            if i in new_ids:
                self.prs.part.drop_rel(sid.rId)
                order.append(new_ids.pop(i))
            else:
                order.append(sid)
        order += [new_ids[k] for k in sorted(new_ids)]  # 원본보다 번호가 큰 장표
        for sid in ids:
            lst.remove(sid)
        for sid in order:
            lst.append(sid)

    @property
    def size_pt(self):
        return self.prs.slide_width / 12700, self.prs.slide_height / 12700

    # ---- 좌표
    def X(self, v):
        return int(self.prs.slide_width * max(0.0, min(100.0, float(v))) / 100)

    def Y(self, v):
        return int(self.prs.slide_height * max(0.0, min(100.0, float(v))) / 100)

    def _color(self, c, default):
        if isinstance(c, str) and HEX.match(c) and self.palette and c.upper().lstrip("#") not in self.palette:
            self.warnings.append(f"가이드 팔레트 밖 색상: {c}")
        return _rgb(c, default)

    # ---- 장표
    def add_slide(self, layout: dict, images: dict[str, Path], crop_fn=None):
        """반환: (slide, {element_id: 실제 사용한 이미지 경로})"""
        layout = normalize(layout)
        slide = self.prs.slides.add_slide(self._blank_layout())
        self.last_sid = self.prs.slides._sldIdLst[-1]
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = self._color(layout.get("background", "#FFFFFF"), "FFFFFF")
        resolved: dict[str, str] = {}
        for i, el in enumerate(layout["elements"]):
            t = el.get("type")
            try:
                fn = getattr(self, f"_el_{t}", None)
                if fn is None:
                    self.warnings.append(f"지원하지 않는 요소 무시: {t}")
                    continue
                shp = fn(slide, el, images, crop_fn, resolved) if t == "image" else fn(slide, el)
                if shp is not None:
                    shp.name = f"el:{el['id']}"
            except Exception as e:  # 요소 하나 실패해도 장표 전체는 살린다
                self.warnings.append(f"{el.get('id')}({t}) 렌더 실패: {e}")
        if layout.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(layout["notes"])
        self.warnings += overflow_warnings(layout, *self.size_pt)
        return slide, resolved

    # ------------------------------------------------------------ 요소들
    def _el_text(self, slide, el):
        from pptx.util import Pt
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        box = slide.shapes.add_textbox(self.X(el["x"]), self.Y(el["y"]), self.X(el["w"]), self.Y(el["h"]))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}.get(
            el.get("valign"), MSO_ANCHOR.TOP)
        size = float(el.get("size", 14))
        if size < self.min_pt:
            self.warnings.append(f"{el['id']}: 글자 크기 {size}pt → {self.min_pt}pt 로 보정")
            size = self.min_pt
        align = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(el.get("align"), PP_ALIGN.LEFT)
        runs = el.get("runs") or [{"text": str(el.get("text", ""))}]
        para = tf.paragraphs[0]

        def style_para(p):
            p.alignment = align
            if el.get("line_spacing"):
                p.line_spacing = float(el["line_spacing"])
        style_para(para)
        for r in runs:
            for j, piece in enumerate(str(r.get("text", "")).split("\n")):
                if j > 0:
                    para = tf.add_paragraph()
                    style_para(para)
                if not piece:
                    continue
                run = para.add_run()
                run.text = piece
                f = run.font
                f.size = Pt(float(r.get("size", size)))
                f.bold = bool(r.get("bold", el.get("bold", False)))
                f.color.rgb = self._color(r.get("color", el.get("color", "#111111")), "111111")
                _set_font(run, el.get("font") or self.font)
        return box

    def _shape(self, slide, el, shape_type):
        from pptx.util import Pt
        shp = slide.shapes.add_shape(shape_type, self.X(el["x"]), self.Y(el["y"]), self.X(el["w"]), self.Y(el["h"]))
        if el.get("fill"):
            shp.fill.solid()
            shp.fill.fore_color.rgb = self._color(el["fill"], "FFFFFF")
        else:
            shp.fill.background()
        if el.get("line"):
            shp.line.color.rgb = self._color(el["line"], "E7E7E4")
            shp.line.width = Pt(float(el.get("line_w", 0.75)))
        else:
            shp.line.fill.background()
        try:
            shp.shadow.inherit = False
        except Exception:
            pass
        if el.get("text"):
            from pptx.enum.text import PP_ALIGN
            tf = shp.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT}.get(el.get("align"), PP_ALIGN.CENTER)
            run = p.add_run()
            run.text = str(el["text"])
            run.font.size = Pt(max(self.min_pt, float(el.get("size", 12))))
            run.font.bold = bool(el.get("bold"))
            run.font.color.rgb = self._color(el.get("color", "#111111"), "111111")
            _set_font(run, self.font)
        return shp

    def _el_rect(self, slide, el):
        from pptx.enum.shapes import MSO_SHAPE
        r = float(el.get("radius", 0) or 0)
        shp = self._shape(slide, el, MSO_SHAPE.ROUNDED_RECTANGLE if r > 0 else MSO_SHAPE.RECTANGLE)
        if r > 0:
            shp.adjustments[0] = max(0.0, min(0.5, r))
        return shp

    def _el_ellipse(self, slide, el):
        from pptx.enum.shapes import MSO_SHAPE
        return self._shape(slide, el, MSO_SHAPE.OVAL)

    def _el_line(self, slide, el):
        from pptx.util import Pt
        from pptx.enum.shapes import MSO_CONNECTOR
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.oxml.ns import qn
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self.X(el["x1"]), self.Y(el["y1"]),
                                       self.X(el["x2"]), self.Y(el["y2"]))
        c.line.color.rgb = self._color(el.get("color", "#111111"), "111111")
        c.line.width = Pt(float(el.get("width", 1)))
        if el.get("dash"):
            c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if el.get("arrow"):
            ln = c.line._get_or_add_ln()
            ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
        return c

    def _el_image(self, slide, el, images, crop_fn, resolved):
        src = str(el.get("source", ""))
        path = images.get(src)
        if path is None and src.startswith(("crop_source", "crop_design")) and crop_fn:
            path = crop_fn(el)
        if not path or not Path(path).exists():
            self.warnings.append(f"{el['id']}: 이미지 없음 ({src}) → 자리표시로 대체")
            return self._el_rect(slide, {**el, "fill": "#F3F3F1", "line": "#E7E7E4",
                                         "text": el.get("label") or "이미지 자리", "size": 10, "color": "#747474"})
        x, y, w, h = self.X(el["x"]), self.Y(el["y"]), self.X(el["w"]), self.Y(el["h"])
        if el.get("fit", "contain") == "contain":
            try:
                from PIL import Image
                with Image.open(path) as im:
                    iw, ih = im.size
                ratio = min(w / iw, h / ih)
                nw, nh = int(iw * ratio), int(ih * ratio)
                x, y, w, h = x + (w - nw) // 2, y + (h - nh) // 2, nw, nh
            except Exception:
                pass
        resolved[el["id"]] = str(path)
        return slide.shapes.add_picture(str(path), x, y, w, h)

    def _chart_common(self, chart):
        from pptx.util import Pt
        chart.font.size = Pt(10)
        chart.font.name = self.font
        chart.font.color.rgb = _rgb("#747474")
        chart.has_legend = False
        chart.has_title = False

    def _el_donut(self, slide, el):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.oxml.ns import qn
        data = CategoryChartData()
        vals = [float(v) for v in el.get("values", [1, 4])]
        data.categories = el.get("labels") or [f"항목{i + 1}" for i in range(len(vals))]
        data.add_series("값", vals)
        gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, self.X(el["x"]), self.Y(el["y"]),
                                    self.X(el["w"]), self.Y(el["h"]), data)
        ch = gf.chart
        self._chart_common(ch)
        plot = ch.plots[0]
        plot.has_data_labels = False
        colors = el.get("colors") or ["#FF6B2C", "#E7E7E4"]
        for i, pt in enumerate(plot.series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = self._color(colors[i % len(colors)], "E7E7E4")
            pt.format.line.fill.background()
        hole = int(float(el.get("hole", 0.68)) * 100)
        for hs in ch._chartSpace.iter(qn("c:holeSize")):
            hs.set("val", str(max(10, min(90, hole))))
        return gf

    def _el_bar(self, slide, el):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
        from pptx.util import Pt
        data = CategoryChartData()
        data.categories = el.get("categories", [])
        data.add_series("값", [float(v) for v in el.get("values", [])])
        kind = XL_CHART_TYPE.BAR_CLUSTERED if el.get("horizontal") else XL_CHART_TYPE.COLUMN_CLUSTERED
        gf = slide.shapes.add_chart(kind, self.X(el["x"]), self.Y(el["y"]), self.X(el["w"]), self.Y(el["h"]), data)
        ch = gf.chart
        self._chart_common(ch)
        va = ch.value_axis
        va.visible = False
        va.has_major_gridlines = False
        ca = ch.category_axis
        ca.format.line.color.rgb = _rgb("#E7E7E4")
        ca.tick_labels.font.size = Pt(float(el.get("label_size", 11)))
        plot = ch.plots[0]
        plot.gap_width = int(el.get("gap", 70))
        plot.vary_by_categories = False
        if el.get("show_values", True):
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(float(el.get("value_size", 12)))
            dl.font.bold = True
            dl.number_format = el.get("number_format", "General")
            dl.number_format_is_linked = False
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
        colors = el.get("colors") or ["#C9C9C6", "#FF6B2C"]
        for i, pt in enumerate(plot.series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = self._color(colors[i % len(colors)], "C9C9C6")
        return gf

    def save(self, out: Path):
        out.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(out))
        return out


def palette_from_guide(guide_md: str) -> set[str]:
    return {m.upper() for m in re.findall(r"#[0-9A-Fa-f]{6}\b", guide_md or "")}
